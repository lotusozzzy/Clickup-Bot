import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
import datetime
import re
import smtplib
import sys
import time
import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from collections import defaultdict

# ==========================================
#              AYARLAR BÖLÜMÜ
# ==========================================
# Hassas bilgiler ortam değişkenlerinden okunur. Bkz: .env.example
def _required_env(name):
    val = os.environ.get(name)
    if not val:
        raise RuntimeError(
            f"Eksik ortam değişkeni: {name}. .env dosyanı kontrol et "
            f"veya 'set -a; source .env; set +a' ile yükle."
        )
    return val


API_TOKEN = _required_env("CLICKUP_API_TOKEN")
GONDEREN_SIFRE = _required_env("SMTP_PASSWORD")

WORKSPACE_ID = os.environ.get("CLICKUP_WORKSPACE_ID", "2598108")
SPACE_ADI = os.environ.get("CLICKUP_SPACE_ADI", "Cariler")
ALAN_ID = os.environ.get("CLICKUP_ALAN_ID", "3c3d2b8c-5346-407f-862b-8ab13f173519")
TARIH_ALAN_ID = os.environ.get("CLICKUP_TARIH_ALAN_ID", "885e079e-dee7-490f-a42a-4d6ea6fa94a8")
SIRKET_ALAN_ID = os.environ.get("CLICKUP_SIRKET_ALAN_ID", "26c369ca-6a0d-48fe-8132-40df9ca5c5b0")
MILAT_TIMESTAMP = int(datetime.datetime(2024, 1, 1).timestamp() * 1000)

GONDEREN_MAIL = os.environ.get("SMTP_FROM", "clickup@solarfame.com")
ALICI_MAIL = os.environ.get("MAIL_TO", "solarframemuhasebe@solarfame.com")
SMTP_SUNUCU = os.environ.get("SMTP_HOST", "smtp.yandex.com.tr")
SMTP_PORT = int(os.environ.get("SMTP_PORT", "465"))

SF_ORANGE = "F26A21"
SF_GRAY = "595959"

# Güvenlik / performans sabitleri
MAX_PAGES_PER_LIST = 200          # Liste başına en fazla 200 sayfa (=20.000 task) dene
REQUEST_TIMEOUT = (10, 25)        # (connect, read) saniye
BASE_SLEEP = 0.7                  # Her istek arasında bekleme
MAX_429_RETRIES = 6               # Tek sayfa için 429 tekrar denemesi

headers = {"Authorization": API_TOKEN, "Content-Type": "application/json"}


def log(msg):
    """stdout buffer'ını anında boşalt - terminalde donmuş gibi görünmesin."""
    print(msg, flush=True)


def get_session():
    """Retry mekanizmalı requests session.

    Not: 429'u burada status_forcelist'e koymuyoruz; kendi elimizle
    Retry-After header'ına saygı duyarak yöneteceğiz. Bu sayede uzun
    backoff'lar nedeniyle script donmuş gibi görünmüyor.
    """
    session = requests.Session()
    retry = Retry(
        total=3,
        connect=3,
        read=3,
        backoff_factor=1.0,
        status_forcelist=[500, 502, 503, 504],
        allowed_methods=["GET"],
        respect_retry_after_header=True,
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry, pool_connections=10, pool_maxsize=10)
    session.mount('http://', adapter)
    session.mount('https://', adapter)
    session.headers.update(headers)
    return session


def safe_get(session, url, attempt_label=""):
    """ClickUp'ı 429-aware şekilde GET et. Hata olursa None döner."""
    for attempt in range(1, MAX_429_RETRIES + 1):
        try:
            r = session.get(url, timeout=REQUEST_TIMEOUT)
        except requests.exceptions.RequestException as e:
            log(f"   ⚠️ Bağlantı hatası ({attempt_label}, deneme {attempt}/{MAX_429_RETRIES}): {e}")
            time.sleep(min(2 ** attempt, 30))
            continue

        if r.status_code == 429:
            # Rate limit. ClickUp X-RateLimit-Reset header'ı (epoch saniye) verir.
            reset = r.headers.get("X-RateLimit-Reset")
            retry_after = r.headers.get("Retry-After")
            wait = 5
            try:
                if retry_after:
                    wait = int(float(retry_after))
                elif reset:
                    wait = max(1, int(float(reset)) - int(time.time()))
            except (ValueError, TypeError):
                pass
            wait = max(1, min(wait, 60))
            log(f"   ⏳ Rate limit (429). {wait}s bekleniyor... ({attempt_label})")
            time.sleep(wait)
            continue

        if r.status_code >= 500:
            log(f"   ⚠️ Sunucu hatası {r.status_code} ({attempt_label}). Tekrar deneniyor.")
            time.sleep(min(2 ** attempt, 30))
            continue

        return r

    log(f"   ❌ Vazgeçildi: {attempt_label}")
    return None


def sayiya_cevir(deger):
    if deger is None:
        return 0.0
    s = str(deger).strip()
    if not s:
        return 0.0
    is_negative = any(char in s for char in ['-', '−', '—'])
    s = re.sub(r'[^\d.,]', '', s)
    if not s:
        return 0.0
    s = s.replace(',', '')
    try:
        return -float(s) if is_negative else float(s)
    except ValueError:
        return 0.0


def fetch_list_tasks(session, l_id, l_name):
    """Tek bir listenin tüm tasklarını sayfa sayfa çek.

    Donmaları engellemek için:
    - last_page alanını ve task sayısını kontrol et
    - Aynı task ID'leri tekrar gelirse döngüyü kır (sonsuz döngü koruması)
    - MAX_PAGES_PER_LIST üst sınırı koy
    """
    seen_task_ids = set()
    all_tasks = []
    page = 0

    while page < MAX_PAGES_PER_LIST:
        t_url = (
            f"https://api.clickup.com/api/v2/list/{l_id}/task"
            f"?include_closed=true&subtasks=true&page={page}"
        )
        label = f"{l_name} sayfa {page}"
        t_res = safe_get(session, t_url, attempt_label=label)

        if t_res is None:
            log(f"   ⚠️ '{l_name}' listesinin {page}. sayfasında vazgeçildi, sonraki listeye geçiliyor.")
            break

        if t_res.status_code != 200:
            log(f"   ⚠️ HTTP {t_res.status_code} ({label}) - liste atlandı.")
            break

        try:
            payload = t_res.json()
        except ValueError:
            log(f"   ⚠️ JSON parse hatası ({label}) - liste atlandı.")
            break

        tasks = payload.get('tasks', []) or []
        if not tasks:
            break

        # Sonsuz döngü koruması: bu sayfadaki tüm task'lar daha önce görüldüyse dur.
        new_ids = {t.get('id') for t in tasks if t.get('id')}
        if new_ids and new_ids.issubset(seen_task_ids):
            log(f"   ℹ️ Tekrar eden sayfa tespit edildi ({label}), pagination kapatıldı.")
            break
        seen_task_ids.update(new_ids)

        all_tasks.extend(tasks)

        # ClickUp last_page alanı dönerse onu da dikkate al.
        if payload.get('last_page') is True:
            break

        # Sayfa 100'den az task döndüyse genelde sondur.
        if len(tasks) < 100:
            break

        page += 1
        time.sleep(BASE_SLEEP)

    if page >= MAX_PAGES_PER_LIST:
        log(f"   ⚠️ '{l_name}' için MAX_PAGES_PER_LIST sınırına ulaşıldı, durduruldu.")

    return all_tasks


def verileri_cek_ve_raporla():
    log("🚀 Başlatılıyor... Veriler çekiliyor.")
    raporlanacak_cariler = []
    session = get_session()

    try:
        r = safe_get(session, f"https://api.clickup.com/api/v2/team/{WORKSPACE_ID}/space", attempt_label="spaces")
        if r is None or r.status_code != 200:
            log("❌ Space listesi alınamadı.")
            return raporlanacak_cariler
        spaces = r.json().get('spaces', [])
        space_id = next((s['id'] for s in spaces if s['name'].lower() == SPACE_ADI.lower()), None)
        if space_id is None:
            log(f"❌ '{SPACE_ADI}' isimli space bulunamadı.")
            return raporlanacak_cariler

        all_lists = []
        r_lists = safe_get(session, f"https://api.clickup.com/api/v2/space/{space_id}/list", attempt_label="space lists")
        if r_lists is not None and r_lists.status_code == 200:
            all_lists.extend(r_lists.json().get('lists', []))

        r_folders = safe_get(session, f"https://api.clickup.com/api/v2/space/{space_id}/folder", attempt_label="folders")
        folders = r_folders.json().get('folders', []) if (r_folders is not None and r_folders.status_code == 200) else []
        for f in folders:
            r_fl = safe_get(session, f"https://api.clickup.com/api/v2/folder/{f['id']}/list", attempt_label=f"folder {f.get('name', f['id'])}")
            if r_fl is not None and r_fl.status_code == 200:
                all_lists.extend(r_fl.json().get('lists', []))

        log(f"📦 Toplam {len(all_lists)} cari listesi taranıyor.\n")
        toplam_islenen_task = 0
        baslangic = time.time()

        for idx, lst in enumerate(all_lists, 1):
            l_id, l_name = lst['id'], lst['name']
            sirket_bakiyeleri = defaultdict(float)

            log(f"[{idx}/{len(all_lists)}] 🔍 {l_name} taranıyor...")
            tasks = fetch_list_tasks(session, l_id, l_name)
            toplam_islenen_task += len(tasks)

            for t in tasks:
                tarih_ok = False
                current_sirket = "Belirtilmemiş"
                alt_deger = 0.0

                custom_fields = t.get('custom_fields', []) or []
                for cf in custom_fields:
                    cf_id = cf.get('id')
                    if cf_id == TARIH_ALAN_ID and 'value' in cf:
                        try:
                            if int(cf['value']) >= MILAT_TIMESTAMP:
                                tarih_ok = True
                        except (ValueError, TypeError):
                            pass
                    if cf_id == SIRKET_ALAN_ID and 'value' in cf:
                        current_sirket = str(cf.get('value', "Belirtilmemiş"))

                if tarih_ok:
                    for cf in custom_fields:
                        if cf.get('id') == ALAN_ID and 'value' in cf:
                            alt_deger = sayiya_cevir(cf['value'])

                    if alt_deger != 0:
                        sirket_bakiyeleri[current_sirket] += alt_deger

            for sirket, bakiye in sirket_bakiyeleri.items():
                if round(bakiye, 2) != 0:
                    raporlanacak_cariler.append({
                        "ad": l_name,
                        "sirket": sirket,
                        "bakiye": round(bakiye, 2),
                    })

            gecen = int(time.time() - baslangic)
            log(f"   ✓ {len(tasks)} task işlendi (toplam {toplam_islenen_task}, {gecen}s).")

    except Exception as e:
        log(f"\n❌ Kritik Hata: {e}")

    return raporlanacak_cariler


def excel_ve_mail(veriler):
    if not veriler:
        log("ℹ️ Raporlanacak veri yok, mail gönderilmiyor.")
        return
    dosya = f"Cari_Rapor_{datetime.datetime.now().strftime('%d_%m_%Y')}.xlsx"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    sirket_gruplari = defaultdict(list)
    for v in veriler:
        sirket_gruplari[v['sirket']].append(v)

    header_fill = PatternFill(start_color=SF_ORANGE, end_color=SF_ORANGE, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    wrap_alignment = Alignment(wrap_text=True, vertical="center")
    thin_border = Border(left=Side(style='thin', color="D3D3D3"),
                         right=Side(style='thin', color="D3D3D3"),
                         top=Side(style='thin', color="D3D3D3"),
                         bottom=Side(style='thin', color="D3D3D3"))

    for sirket_adi, cari_listesi in sirket_gruplari.items():
        safe_sheet_name = re.sub(r'[\\/*?:\[\]]', '', sirket_adi)[:31]
        if not safe_sheet_name:
            safe_sheet_name = "Bilinmeyen"
        ws = wb.create_sheet(title=safe_sheet_name)

        basliklar = ["Cari Adı", "Ait Olduğu Şirket", "Bakiye (TL)"]
        ws.append(basliklar)

        header_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # Başlıkların iki satıra sararken kesilmemesi için yeterli yükseklik
        ws.row_dimensions[1].height = 30

        for cari in cari_listesi:
            ws.append([cari['ad'], cari['sirket'], cari['bakiye']])

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=3):
            for cell in row:
                cell.alignment = wrap_alignment
                cell.border = thin_border
                if cell.column == 3:
                    cell.number_format = '#,##0.00'

        # Sütun genişlikleri: veriye göre otomatik, ama başlıktaki en uzun
        # kelime de sığsın diye minimum garanti ediliyor. Üst sınır 40.
        for col_idx in range(1, len(basliklar) + 1):
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            max_data_len = 0
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                                    min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value is None:
                        continue
                    if isinstance(cell.value, (int, float)):
                        text = f"{cell.value:,.2f}"
                    else:
                        text = str(cell.value)
                    if len(text) > max_data_len:
                        max_data_len = len(text)

            header_text = basliklar[col_idx - 1]
            min_for_header = max(
                (len(w) for w in header_text.split()), default=len(header_text)
            )
            width = max(min_for_header + 2, max_data_len + 3)
            ws.column_dimensions[col_letter].width = min(width, 40)

        tablo_ref = f"A1:C{ws.max_row}"
        tablo_ismi = "Tbl_" + re.sub(r'[^a-zA-Z0-9]', '', sirket_adi)
        if tablo_ismi == "Tbl_":
            tablo_ismi = "Tbl_Sirket_" + str(id(sirket_adi))

        tab = Table(displayName=tablo_ismi[:255], ref=tablo_ref)
        style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        tab.tableStyleInfo = style
        ws.add_table(tab)

    wb.save(dosya)

    msg = MIMEMultipart()
    msg['From'], msg['To'], msg['Subject'] = GONDEREN_MAIL, ALICI_MAIL, "Haftalık Bakiye Raporu"

    html_govde = f"""
    <html>
      <body style="font-family: Arial, sans-serif; color: #{SF_GRAY};">
        <h2 style="color: #{SF_ORANGE};">Haftalık Bakiye Raporu</h2>
        <p>Açık cariler, ait oldukları şirketlere göre sekmelere ayrılarak ekteki Excel dosyasında sunulmuştur.</p>
        <p><i>Not: Başlıklardaki ok işaretlerine tıklayarak bakiyeleri büyükten küçüğe sıralayabilir veya belirli bir cariyi filtreleyebilirsiniz.</i></p>
        <br>
        <p>İyi çalışmalar.</p>
      </body>
    </html>
    """
    msg.attach(MIMEText(html_govde, 'html'))

    with open(dosya, "rb") as f:
        part = MIMEBase("application", "octet-stream")
        part.set_payload(f.read())
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", f"attachment; filename={dosya}")
    msg.attach(part)

    try:
        with smtplib.SMTP_SSL(SMTP_SUNUCU, SMTP_PORT, timeout=20) as s:
            s.login(GONDEREN_MAIL, GONDEREN_SIFRE)
            s.send_message(msg)
        log("✨ Şirket bazlı, dinamik tablolu kurumsal rapor gönderildi!")
    except Exception as e:
        log(f"❌ Mail Hatası: {e}")
    finally:
        if os.path.exists(dosya):
            os.remove(dosya)


if __name__ == "__main__":
    sonuc = verileri_cek_ve_raporla()
    excel_ve_mail(sonuc)
