"""Coalexus Pipeline Update — günlük rapor.

Coalexus → CRM → "Müşteri Listesi" (list_id=901523815423) üzerinde, bir
önceki BAŞARILI çalışmadan beri GÜNCELLENEN veya OLUŞTURULAN müşteri
(satış-hattı) task'larını Excel olarak Coalexus satış ekibine maille gönderir.

Tasarım kararları (ADIM 1 keşfiyle doğrulandı):
  - View 2f96w-420315'in tek filtresi "dateUpdated lastXDays x=2". Bu
    pencere uzun tatil sonrası "eski-ama-yeni-güncellenmiş" task'ları
    DÜŞÜRÜR → boşluk. Bu yüzden fetch view/task endpoint'inden DEĞİL,
    doğrudan kaynak liste endpoint'inden yapılır; pencere tamamen
    kendi watermark'ımızla (date_updated/date_created >= W) belirlenir.
  - Watermark (last_successful_run_ms) SADECE mail başarıyla gönderilince
    run-başı zamanına (T0) ilerler. Herhangi bir hata/exception/SMTP fail
    → state'e dokunulmaz → sonraki run aynı pencereden devam eder (boşluk yok).
  - Resmi tatilde (TR, dini bayramlar dahil) mail atılmaz VE watermark
    ilerlemez → değişiklikler ertesi iş gününe taşınır. Hafta sonu zaten
    cron (Pzt-Cum) ile tetiklenmez.
  - "Değişiklik yok" günü: 0 task olsa bile kısa bilgi maili atılır ve
    watermark ilerler.

Custom field DEĞERLERİ toplu liste yanıtında geliyor (ek task çağrısı yok);
sadece "Son Yorum"/"Yorum Sayısı" için filtrelenmiş (küçük) alt-kümeye
per-task /comment çağrısı yapılır.

Hassas ayarlar (API token, mail şifresi) clickup_bot modülünden import edilir;
o da değerleri os.environ'dan okur. Reuse: clickup_bot (oturum/fetch/SMTP
sabitleri) + clickup_due_report (fmt_*, _add_sheet, yorum helper'ları).
"""

import datetime
import json
import os
import smtplib
import sys
import time
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import holidays
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

try:
    from zoneinfo import ZoneInfo
except ImportError:  # py<3.9 fallback (sunucuda 3.9+ bekleniyor)
    ZoneInfo = None

from clickup_bot import (
    BASE_SLEEP,
    GONDEREN_MAIL,
    GONDEREN_SIFRE,
    KEEP_EXCEL_ON_DISK,
    MAX_PAGES_PER_LIST,
    SF_GRAY,
    SF_ORANGE,
    SMTP_PORT,
    SMTP_SUNUCU,
    get_session,
    log,
    safe_get,
)
from clickup_due_report import (
    _add_sheet,
    _extract_comment_text,
    fetch_all_comments,
    fmt_date,
    fmt_datetime,
)

# ---------------------------------------------------------------------------
# Sabitler
# ---------------------------------------------------------------------------
LIST_ID = os.environ.get("COALEXUS_LIST_ID", "901523815423")
LIST_NAME = "Müşteri Listesi"
REPORT_NAME = "Coalexus Pipeline Update"

# "Açıklama" custom field (type=text) — get_custom_fields ile doğrulandı.
DESC_CF_ID = "397aee00-91e4-4d18-8f1d-13759b7c6255"

RECIPIENTS = [
    "ozan.koroglu@coalexus.com",
    "efe.soykan@coalexus.com",
    "fatih.yazici@coalexus.com",
]

STATE_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "coalexus_pipeline_state.json"
)

# İlk run'da (state yoksa) bakılacak geriye-dönük pencere.
BOOTSTRAP_LOOKBACK_MS = 24 * 3600 * 1000

# ClickUp /comment endpoint tek istekte en yeni ~25 yorumu döner. Yorum
# sayısı bu sınıra dayanırsa "25+" gösterilir (Ders 2: 25-yorum limiti).
COMMENT_PAGE_CAP = 25

# Hücrelerin şişmemesi için uzun metin kırpma sınırı.
_TEXT_TRUNCATE = 500

COLUMNS = [
    "Müşteri/Task Adı",
    "Açıklama",
    "Son Yorum",
    "Bitiş Tarihi",
    "Son Güncelleme",
    "Öncelik",
    "Yorum Sayısı",
]


# ---------------------------------------------------------------------------
# Test/override env'leri
# ---------------------------------------------------------------------------
def _env_truthy(name):
    """Env değeri truthy mi (KEEP_EXCEL_ON_DISK ile aynı kural)."""
    return os.environ.get(name, "").strip().lower() in ("true", "1", "yes", "on")


def _parse_recipients(raw):
    """Virgülle ayrılmış mail listesini temizleyip liste döndür. Boş → []."""
    if not raw:
        return []
    return [addr.strip() for addr in raw.split(",") if addr.strip()]


# ---------------------------------------------------------------------------
# Zaman / state
# ---------------------------------------------------------------------------
def now_ms():
    """Şu anki epoch ms (timezone-agnostik; pencere hesapları hep ms ile)."""
    return int(time.time() * 1000)


def today_istanbul():
    """Europe/Istanbul yerel tarihini döndür (tatil kontrolü için).

    zoneinfo varsa onu kullan; yoksa sunucu zaten Europe/Istanbul olduğu
    için naive now() yeterli (Ders 3: TZ karışıklığından kaçın).
    """
    if ZoneInfo is not None:
        try:
            return datetime.datetime.now(ZoneInfo("Europe/Istanbul")).date()
        except Exception:  # tzdata yoksa fallback
            pass
    return datetime.datetime.now().date()


def turkish_holiday_name(today):
    """today bir TR resmi tatiliyse tatil adını, değilse None döndür.

    holidays.country_holidays('TR', ...) dini bayramları (Ramazan/Kurban)
    da kapsar (holidays>=0.83 ile doğrulandı).
    """
    tr = holidays.country_holidays("TR", years=today.year)
    return tr.get(today)


def load_state():
    """State dosyasını oku. Yok/bozuksa None döndür (sıfırdan başla)."""
    if not os.path.exists(STATE_FILE):
        return None
    try:
        with open(STATE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except (IOError, json.JSONDecodeError) as e:
        log(f"⚠️ State okunamadı ({e}), watermark sıfırlanıyor (bootstrap).")
        return None


def compute_watermark(state, t0):
    """(W, is_bootstrap) döndür.

    W = state.last_successful_run_ms; yoksa bootstrap = t0 - 24h.
    """
    if state and isinstance(state.get("last_successful_run_ms"), (int, float)):
        return int(state["last_successful_run_ms"]), False
    return t0 - BOOTSTRAP_LOOKBACK_MS, True


def save_state(last_successful_run_ms, status):
    """Watermark'ı atomik (tmp + fsync + os.replace) yaz.

    SADECE mail başarıyla gönderildikten SONRA çağrılır. Yarıda kesintide
    bozuk JSON kalmasın diye .tmp'e yazıp rename edilir (Ders 7).
    """
    payload = {
        "schema": 1,
        "last_successful_run_ms": int(last_successful_run_ms),
        "saved_at": datetime.datetime.now().isoformat(),
        "last_status": status,
    }
    tmp_path = STATE_FILE + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp_path, STATE_FILE)


# ---------------------------------------------------------------------------
# Kaynak liste fetch (BOŞLUKSUZLUK için sıkı — eksikse fırlatır)
# ---------------------------------------------------------------------------
def fetch_all_tasks_strict(session):
    """Kaynak listenin TÜM task'larını (include_closed) sayfalı çek.

    clickup_bot.fetch_list_tasks bu rapor için KASTEN reuse EDİLMEDİ: o,
    herhangi bir sayfada API hatası/non-200/JSON hatası görünce sessizce
    `break` edip KISMİ (veya boş) liste döndürür ve asla exception fırlatmaz.
    Bu, watermark'ın eksik bir pencere üzerinde T0'a ilerlemesine ve
    güncellenmiş task'ların KALICI olarak kaçırılmasına yol açar
    (gap-free garantisinin ihlali). Burada her başarısızlık RuntimeError
    fırlatır → __main__ guard'ı sys.exit(1) yapar → save_state'e ulaşılmaz →
    watermark ilerlemez → sonraki run aynı pencereden devam eder.

    Rate-limit/retry primitifi (safe_get) yine clickup_bot'tan reuse edilir.
    """
    tasks = []
    seen = set()
    page = 0
    while page < MAX_PAGES_PER_LIST:
        url = (
            f"https://api.clickup.com/api/v2/list/{LIST_ID}/task"
            f"?include_closed=true&subtasks=true&page={page}"
        )
        r = safe_get(session, url, attempt_label=f"{LIST_NAME} sayfa {page}")
        if r is None:
            raise RuntimeError(
                f"Liste çekme başarısız (sayfa {page}): safe_get None döndü — "
                f"watermark ilerletilmeyecek."
            )
        if r.status_code != 200:
            raise RuntimeError(
                f"Liste çekme HTTP {r.status_code} (sayfa {page})."
            )
        try:
            payload = r.json()
        except ValueError as e:
            raise RuntimeError(f"Liste çekme JSON parse hatası (sayfa {page}): {e}")

        page_tasks = payload.get("tasks", []) or []
        if not page_tasks:
            break
        new_ids = {t.get("id") for t in page_tasks if t.get("id")}
        if new_ids and new_ids.issubset(seen):
            break  # tekrar eden sayfa — pagination sonu
        seen.update(new_ids)
        tasks.extend(page_tasks)
        if payload.get("last_page") is True:
            break
        if len(page_tasks) < 100:
            break
        page += 1
        time.sleep(BASE_SLEEP)

    if page >= MAX_PAGES_PER_LIST:
        raise RuntimeError(
            f"MAX_PAGES_PER_LIST ({MAX_PAGES_PER_LIST}) aşıldı — liste eksik "
            f"çekilmiş olabilir, watermark ilerletilmiyor."
        )
    return tasks


# ---------------------------------------------------------------------------
# Pencere filtresi + alan çıkarımı
# ---------------------------------------------------------------------------
def _to_ms(value):
    """ClickUp ms-string/int alanını int'e çevir. Parse edilemezse None."""
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (ValueError, TypeError):
        return None


def task_in_window(task, watermark_ms):
    """Task pencerede mi: date_updated >= W VEYA date_created >= W.

    İkisi de yoksa (parse edilemezse) dışarıda bırakılır. Bu fonksiyon
    raporun boşluksuzluğunun çekirdeğidir — saf ve test edilebilir tutulur.
    """
    updated = _to_ms(task.get("date_updated"))
    created = _to_ms(task.get("date_created"))
    if updated is not None and updated >= watermark_ms:
        return True
    if created is not None and created >= watermark_ms:
        return True
    return False


def _truncate(text):
    if text is None:
        return ""
    text = str(text)
    if len(text) > _TEXT_TRUNCATE:
        return text[: _TEXT_TRUNCATE - 1] + "…"
    return text


def get_description(task):
    """'Açıklama' (text) custom field'ının düz metin değerini döndür."""
    for cf in task.get("custom_fields", []) or []:
        if cf.get("id") == DESC_CF_ID:
            return _truncate(cf.get("value") or "")
    return ""


def get_priority(task):
    """ClickUp priority dict'inden etiket döndür ('urgent'/'high'/...).

    priority null olabilir veya {'priority': 'urgent', ...} biçiminde gelir
    (gerçek task üzerinde doğrulandı — string DEĞİL, dict).
    """
    p = task.get("priority")
    if isinstance(p, dict):
        return p.get("priority") or ""
    if isinstance(p, str):  # savunmacı
        return p
    return ""


def fetch_comment_summary(session, task_id):
    """(son_yorum_metni, yorum_sayisi_gosterimi) döndür.

    fetch_all_comments newest-first sıralı döner (kendi içinde reverse sort);
    comments[0] = en yeni yorum. Tek sayfa (~25) çekilir; sayı bu sınıra
    dayanırsa '25+' gösterilir (Ders 2).
    """
    comments = fetch_all_comments(session, task_id)
    if not comments:
        return "", "0"
    latest = _truncate(_extract_comment_text(comments[0]))
    n = len(comments)
    count_str = f"{COMMENT_PAGE_CAP}+" if n >= COMMENT_PAGE_CAP else str(n)
    return latest, count_str


def build_rows(session, tasks):
    """Filtrelenmiş task'lardan Excel satırlarını (kolon sırasıyla) üret.

    En son güncellenen en üstte olacak şekilde sıralanır.
    """
    ordered = sorted(
        tasks, key=lambda t: _to_ms(t.get("date_updated")) or 0, reverse=True
    )
    rows = []
    for t in ordered:
        latest_comment, comment_count = fetch_comment_summary(session, t.get("id"))
        rows.append([
            t.get("name", ""),
            get_description(t),
            latest_comment,
            fmt_date(t.get("due_date")),
            fmt_datetime(t.get("date_updated")),
            get_priority(t),
            comment_count,
        ])
        time.sleep(BASE_SLEEP)
    return rows


# ---------------------------------------------------------------------------
# Excel + Mail
# ---------------------------------------------------------------------------
def write_excel(rows):
    """Tek sayfalı Excel üret, dosya adını döndür."""
    dosya = f"{REPORT_NAME}_{datetime.datetime.now().strftime('%Y-%m-%d')}.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color=SF_ORANGE, end_color=SF_ORANGE, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    wrap_alignment = Alignment(wrap_text=True, vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )
    _add_sheet(
        wb, "Pipeline", COLUMNS, rows, "Tbl_Pipeline",
        header_fill, header_font, wrap_alignment, thin_border,
    )
    wb.save(dosya)
    return dosya


def _build_mail(dosya, row_count, is_bootstrap, recipients):
    konu_tarih = datetime.datetime.now().strftime("%d.%m.%Y")
    msg = MIMEMultipart()
    msg["From"] = GONDEREN_MAIL
    msg["To"] = ", ".join(recipients)
    msg["Subject"] = f"{REPORT_NAME} — {konu_tarih}"

    bootstrap_note = ""
    if is_bootstrap:
        bootstrap_note = (
            "<p style='color:#888;font-size:12px;'>İlk çalıştırma: pencere "
            "son 24 saat olarak başlatıldı. Sonraki raporlar son başarılı "
            "çalışmadan beri olan değişiklikleri kapsayacak.</p>"
        )

    if row_count > 0:
        body = f"""
        <html><body style="font-family: Arial, sans-serif; color: #{SF_GRAY};">
          <h2 style="color: #{SF_ORANGE};">{REPORT_NAME}</h2>
          <p>Son başarılı çalışmadan bu yana <b>{row_count}</b> müşteri kaydı
             güncellendi veya oluşturuldu. Detaylar ekteki Excel dosyasındadır.</p>
          {bootstrap_note}
          <p>İyi çalışmalar.</p>
        </body></html>
        """
    else:
        body = f"""
        <html><body style="font-family: Arial, sans-serif; color: #{SF_GRAY};">
          <h2 style="color: #{SF_ORANGE};">{REPORT_NAME}</h2>
          <p>Bugün son başarılı çalışmadan bu yana güncellenen veya oluşturulan
             müşteri kaydı bulunmuyor.</p>
          {bootstrap_note}
          <p>İyi çalışmalar.</p>
        </body></html>
        """
    msg.attach(MIMEText(body, "html"))

    if dosya:
        with open(dosya, "rb") as f:
            part = MIMEBase("application", "octet-stream")
            part.set_payload(f.read())
        encoders.encode_base64(part)
        # filename boşluk içeriyor ("Coalexus Pipeline Update_...xlsx"); keyword
        # form değeri RFC'ye uygun şekilde tırnaklar (f-string'le tırnaksız
        # kalır ve katı istemciler ilk boşlukta keser).
        part.add_header("Content-Disposition", "attachment", filename=dosya)
        msg.attach(part)
    return msg


def send_report_mail(dosya, row_count, is_bootstrap, recipients=None):
    """Raporu verilen alıcılara gönder (None → sabit RECIPIENTS).
    Başarısızlıkta exception fırlatır (böylece watermark İLERLEMEZ).
    send_mail/excel_ve_mail rapora-özel olduğu için import edilmedi; SMTP_SSL
    kalıbı burada yeniden kuruldu.
    """
    recipients = recipients or RECIPIENTS
    msg = _build_mail(dosya, row_count, is_bootstrap, recipients)
    with smtplib.SMTP_SSL(SMTP_SUNUCU, SMTP_PORT, timeout=30) as s:
        s.login(GONDEREN_MAIL, GONDEREN_SIFRE)
        # send_message TÜM alıcılar reddedilirse SMTPRecipientsRefused fırlatır,
        # ama KISMİ ret (1-2/3) sessizce reddedilenlerin dict'ini döndürür.
        # Stateful watermark'ta sessiz kısmi ret = kalıcı boşluk → fırlat.
        refused = s.send_message(msg)
    if refused:
        raise smtplib.SMTPRecipientsRefused(refused)
    log(f"✨ Mail gönderildi → {', '.join(recipients)}")


# ---------------------------------------------------------------------------
# Ana akış
# ---------------------------------------------------------------------------
def main():
    t0 = now_ms()
    log(f"🚀 {REPORT_NAME} başlatılıyor — {datetime.datetime.now()}")

    # 1) Tatil kontrolü — tatilse mail YOK, watermark İLERLEMEZ, çık.
    today = today_istanbul()
    holiday_name = turkish_holiday_name(today)
    if holiday_name:
        log(f"🎌 Bugün resmi tatil ({holiday_name}). Mail atılmadı, "
            f"watermark ilerletilmedi — değişiklikler ertesi iş gününe taşınacak.")
        return

    # 2) Watermark
    state = load_state()
    watermark, is_bootstrap = compute_watermark(state, t0)
    log(f"⏱️  Watermark W={watermark} ({fmt_datetime(watermark)}), "
        f"bootstrap={is_bootstrap}, T0={t0} ({fmt_datetime(t0)}).")

    # 3) Kaynak listeyi çek (view/task DEĞİL — pencere kendi watermark'ımız).
    #    Eksik/başarısız fetch'te fırlatır → watermark ilerlemez (boşluk yok).
    session = get_session()
    tasks = fetch_all_tasks_strict(session)
    log(f"📦 '{LIST_NAME}' içinden {len(tasks)} task çekildi (include_closed).")

    # 4) Pencere filtresi: date_updated >= W VEYA date_created >= W
    filtered = [t for t in tasks if task_in_window(t, watermark)]
    log(f"🔎 Pencere içi (güncellenen/oluşturulan): {len(filtered)} task.")

    # Test/override modu: DRY_RUN (mail atma) veya RECIPIENT_OVERRIDE (alıcıyı
    # değiştir) set ise bu bir TEST'tir → watermark İLERLEMEZ (mail başarılı
    # olsa bile). Watermark yalnız TEMİZ gerçek run'da (ikisi de yok) +
    # mail başarısında ilerler; böylece testler ilk gerçek cron penceresini
    # etkilemez (state ilk kez gerçek cron run'ında yazılır).
    dry_run = _env_truthy("DRY_RUN")
    override = _parse_recipients(os.environ.get("RECIPIENT_OVERRIDE"))
    test_mode = dry_run or bool(override)
    recipients = override or RECIPIENTS
    # DRY_RUN'da Excel incelensin diye diskte tutulur.
    keep_excel = KEEP_EXCEL_ON_DISK or dry_run

    # 5) Satırlar + Excel
    rows = build_rows(session, filtered)
    dosya = write_excel(rows) if rows else None

    # 6-7) Mail + watermark. save_state SADECE send başarılıysa VE temiz gerçek
    #      run'da çağrılır → watermark yalnız başarıda ilerler. Excel temizliği
    #      finally'de — send fail etse de dosya diskte kalmaz (DRY_RUN hariç).
    try:
        if dry_run:
            log(f"🧪 DRY_RUN: mail atlandı, watermark ilerletilmedi. "
                f"(Excel: {dosya or '(satır yok)'})")
        else:
            if test_mode:
                log(f"🧪 RECIPIENT_OVERRIDE aktif → sadece "
                    f"{', '.join(recipients)} (test, watermark ilerletilmeyecek).")
            send_report_mail(dosya, len(rows), is_bootstrap, recipients)
            if test_mode:
                log("🧪 Test gönderimi tamam; watermark ilerletilmedi.")
            else:
                save_state(t0, "sent" if rows else "no_change")
                log(f"✅ Watermark T0={t0} kaydedildi "
                    f"(status={'sent' if rows else 'no_change'}).")
    finally:
        if dosya and not keep_excel:
            try:
                os.remove(dosya)
            except OSError as e:
                log(f"⚠️ Excel silinemedi ({e}).")
        elif dosya:
            log(f"📁 Excel diskte tutuluyor → '{dosya}' "
                f"(KEEP_EXCEL_ON_DISK veya DRY_RUN).")

    log("✅ İşlem tamamlandı.")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        # Watermark zaten ilerlemedi (save_state'e ulaşılmadı). Cron log'una
        # düş ve non-zero exit ile çık — sonraki run aynı pencereden devam eder.
        log(f"❌ Kritik hata: {e}")
        sys.exit(1)
