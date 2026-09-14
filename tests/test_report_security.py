"""Report regressions using synthetic data and a mocked mail transport.

Run with: python3 -m unittest discover -s tests -v
"""

import io
import os
from pathlib import Path
import tempfile
import unittest
from unittest import mock
import warnings
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

with mock.patch.dict(os.environ, {
    "CLICKUP_API_TOKEN": "test-token",
    "SMTP_PASSWORD": "test-password",
    "SMTP_HOST": "smtp.example.test",
    "SMTP_FROM": "reports@example.test",
    "MAIL_TO": "accounts@example.test",
    "MAIL_TO_DAILY": "management@example.test",
}):
    import clickup_bot as bot
    import clickup_due_report as daily


class ReportSecurityTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        previous = Path.cwd()
        os.chdir(self.directory.name)
        self.addCleanup(os.chdir, previous)
        self.network = mock.patch(
            "socket.create_connection", side_effect=AssertionError("Network forbidden")
        )
        self.network.start()
        self.addCleanup(self.network.stop)
        self.smtp_patch = mock.patch.object(bot.smtplib, "SMTP_SSL")
        self.smtp = self.smtp_patch.start()
        self.addCleanup(self.smtp_patch.stop)
        self.connection = self.smtp.return_value.__enter__.return_value
        log_patch = mock.patch.object(bot, "log")
        log_patch.start()
        self.addCleanup(log_patch.stop)

    def read_workbook(self, content):
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            for name in archive.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    root = ET.fromstring(archive.read(name))
                    self.assertEqual(root.findall(".//s:f", namespace), [], name)
        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def weekly_workbook(self, rows, skipped=None):
        self.connection.send_message.reset_mock()
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            bot.excel_ve_mail(rows, skipped)
        self.connection.send_message.assert_called_once()
        message = self.connection.send_message.call_args.args[0]
        self.assertEqual(message["To"], bot.ALICI_MAIL)
        attachments = [p for p in message.walk() if p.get_content_disposition() == "attachment"]
        self.assertEqual(len(attachments), 1)
        self.assertEqual(list(Path.cwd().glob("*.xlsx")), [])
        return self.read_workbook(attachments[0].get_payload(decode=True))

    def daily_workbook(self, with_comment, no_comment, removed, skipped=None):
        path = Path(daily.write_excel(with_comment, no_comment, removed, skipped))
        return self.read_workbook(path.read_bytes())

    def assert_text(self, cell, expected):
        self.assertEqual(cell.value, expected)
        self.assertEqual(cell.data_type, "s")

    def test_weekly_text_stays_literal_and_balances_stay_numeric(self):
        texts = ["=1+1", '=HYPERLINK("https://example.invalid","text")',
                 "#N/A", "+1", "-1", "@SUM(1,2)", "\n=1+1", "  =1+1",
                 "İstanbul — açıklama\nikinci satır"]
        company = '="Şirket"'
        rows = [{"ad": text, "sirket": company, "bakiye": -1234.56} for text in texts]
        workbook = self.weekly_workbook(rows)
        sheet = workbook.worksheets[0]
        self.assertEqual([c.value for c in sheet[1]], ["Cari Adı", "Ait Olduğu Şirket", "Bakiye (TL)"])
        for number, text in enumerate(texts, 2):
            self.assert_text(sheet.cell(number, 1), text)
            self.assert_text(sheet.cell(number, 2), company)
            balance = sheet.cell(number, 3)
            self.assertEqual(balance.value, -1234.56)
            self.assertEqual(balance.data_type, "n")
            self.assertEqual(balance.number_format, "#,##0.00")
            self.assertTrue(balance.alignment.wrap_text)
        table = next(iter(sheet.tables.values()))
        self.assertEqual(table.tableStyleInfo.name, "TableStyleLight1")

    def test_weekly_skip_only_report_keeps_text(self):
        workbook = self.weekly_workbook([], [{"ad": "=1+1", "sebep": "#N/A"}])
        sheet = workbook["Atlanan Listeler"]
        self.assert_text(sheet["A2"], "=1+1")
        self.assert_text(sheet["B2"], "#N/A")

    def test_colliding_company_labels_all_reach_the_attachment(self):
        labels = ["A-B", "AB", "Case", "case", "Ü", "Ş",
                  "A" * 260 + "X", "A" * 260 + "Y", "Atlanan_Listeler", "[]"]
        rows = [{"ad": "normal", "sirket": label, "bakiye": 10.25} for label in labels]
        workbook = self.weekly_workbook(rows, [{"ad": "skipped", "sebep": "timeout"}])
        self.assertEqual(len(workbook.worksheets), len(labels) + 1)
        self.assertEqual([s["B2"].value for s in workbook.worksheets[:-1]], labels)
        names = [name.lower() for sheet in workbook for name in sheet.tables]
        self.assertEqual(len(names), len(labels) + 1)
        self.assertEqual(len(names), len(set(names)))

    def test_all_daily_categories_and_skipped_rows_stay_literal(self):
        row = {"name": "=1+1", "space_name": '="space"', "list_name": "=2+2",
               "old_date": "01.09.2026", "new_date": "02.09.2026", "changer": "#N/A",
               "last_comment_by": "=3+3", "last_comment_text": "=4+4",
               "url": "https://example.invalid/task"}
        workbook = self.daily_workbook([row], [row], [row],
                                      [{"ad": "=5+5", "sebep": "#N/A"}])
        for sheet in workbook.worksheets[:3]:
            self.assert_text(sheet["A2"], row["name"])
            self.assert_text(sheet["B2"], row["space_name"])
            self.assert_text(sheet["C2"], row["list_name"])
            self.assert_text(sheet["D2"], row["old_date"])
        self.assert_text(workbook.worksheets[0]["F2"], row["changer"])
        self.assert_text(workbook.worksheets[0]["G2"], row["last_comment_by"])
        self.assert_text(workbook.worksheets[0]["H2"], row["last_comment_text"])
        self.assert_text(workbook.worksheets[3]["A2"], "=5+5")
        self.assert_text(workbook.worksheets[3]["B2"], "#N/A")

    def test_snapshot_fallback_and_rich_comment_remain_literal(self):
        previous = {"tasks": {"t": {"name": "=1+1", "due_date": "1000"}}}
        current = {"t": {"name": "", "due_date": "2000", "list_name": "normal",
                         "space_name": "normal", "url": "https://example.invalid/task"}}
        comment = daily._comment_text({"comment": [{"text": "="}, {"text": "2+2"}]})
        with mock.patch.object(daily, "get_due_date_history", return_value=[]), \
                mock.patch.object(daily, "get_last_comment", return_value={
                    "user_id": "collaborator", "username": "person", "text": comment}), \
                mock.patch.object(daily.time, "sleep"):
            groups = daily.diff_snapshots(previous, current, mock.Mock(), "owner")
        workbook = self.daily_workbook(*groups)
        self.assert_text(workbook.worksheets[0]["A2"], "=1+1")
        self.assert_text(workbook.worksheets[0]["H2"], "=2+2")

    def test_empty_daily_categories_keep_their_placeholder_rows(self):
        workbook = self.daily_workbook([], [], [])
        self.assertEqual(len(workbook.worksheets), 3)
        for sheet in workbook:
            self.assert_text(sheet["A2"], "(Bu kategoride değişiklik yok)")
            self.assertEqual(len(sheet.tables), 1)
            self.assertEqual(sheet["A1"].fill.fgColor.rgb[-6:], bot.SF_ORANGE)

    def test_weekly_no_data_does_not_send_mail(self):
        bot.excel_ve_mail([])
        self.smtp.assert_not_called()


if __name__ == "__main__":
    unittest.main()
