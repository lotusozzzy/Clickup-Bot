"""Coalexus Pipeline Update raporu testleri.

Odak (gerçek akışı yakalayan, vacuous olmayan testler):
  - Pencere filtresi: date_updated >= W VEYA date_created >= W  (boşluksuzluk çekirdeği)
  - Watermark SADECE mail başarılıysa ilerler; hata/tatil → ilerlemez
  - Bootstrap (state yok → now-24h)
  - Atomik state yazımı
  - Yorum özeti (newest-first → comments[0])

MUTATION CHECK (manuel doğrulandı, README aşağıda):
  - task_in_window'daki `or` → `and` yapılırsa test_window_created_only FAIL eder.
  - save_state çağrısı send_report_mail'den ÖNCEYE alınırsa (veya finally'ye)
    test_watermark_not_advanced_on_send_failure FAIL eder.
"""

import json
import os
import sys

import pytest


def _import_module(monkeypatch, tmp_path):
    """coalexus_pipeline_report'u taze import edip STATE_FILE'ı tmp'e patch et."""
    monkeypatch.setenv("CLICKUP_API_TOKEN", "test-token-not-real")
    monkeypatch.setenv("SMTP_PASSWORD", "test-smtp-pw")
    monkeypatch.setenv("CLICKUP_WORKSPACE_ID", "9999999")

    for mod_name in list(sys.modules):
        if (
            mod_name in ("clickup_bot", "clickup_due_report", "coalexus_pipeline_report")
            or mod_name.startswith("webhook")
        ):
            del sys.modules[mod_name]

    import coalexus_pipeline_report as mod

    state = tmp_path / "coalexus_pipeline_state.json"
    monkeypatch.setattr(mod, "STATE_FILE", str(state))
    return mod, str(state)


# ---------------------------------------------------------------------------
# Pencere filtresi
# ---------------------------------------------------------------------------
def test_window_updated_included(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    W = 1_000_000
    task = {"date_updated": str(W + 5), "date_created": str(W - 999)}
    assert mod.task_in_window(task, W) is True


def test_window_created_only(monkeypatch, tmp_path):
    """date_updated < W AMA date_created >= W → DAHİL.

    MUTATION: task_in_window'da `or` → `and` yapılırsa bu test FAIL eder.
    """
    mod, _ = _import_module(monkeypatch, tmp_path)
    W = 1_000_000
    task = {"date_updated": str(W - 10), "date_created": str(W + 1)}
    assert mod.task_in_window(task, W) is True


def test_window_both_below_excluded(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    W = 1_000_000
    task = {"date_updated": str(W - 1), "date_created": str(W - 50)}
    assert mod.task_in_window(task, W) is False


def test_window_boundary_inclusive(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    W = 1_000_000
    assert mod.task_in_window({"date_updated": str(W), "date_created": "0"}, W) is True


def test_window_unparseable_excluded(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    assert mod.task_in_window({"date_updated": None, "date_created": ""}, 1) is False
    assert mod.task_in_window({"date_updated": "xx", "date_created": "yy"}, 1) is False


# ---------------------------------------------------------------------------
# Watermark / bootstrap / state
# ---------------------------------------------------------------------------
def test_bootstrap_when_no_state(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    t0 = 5_000_000_000
    W, is_bootstrap = mod.compute_watermark(None, t0)
    assert is_bootstrap is True
    assert W == t0 - mod.BOOTSTRAP_LOOKBACK_MS


def test_watermark_from_state(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    t0 = 5_000_000_000
    W, is_bootstrap = mod.compute_watermark({"last_successful_run_ms": 123456}, t0)
    assert is_bootstrap is False
    assert W == 123456


def test_save_state_atomic(monkeypatch, tmp_path):
    mod, state_path = _import_module(monkeypatch, tmp_path)
    mod.save_state(777, "sent")
    assert os.path.exists(state_path)
    assert not os.path.exists(state_path + ".tmp"), "tmp os.replace sonrası kalmamalı"
    with open(state_path, encoding="utf-8") as f:
        data = json.load(f)
    assert data["last_successful_run_ms"] == 777
    assert data["last_status"] == "sent"


def test_load_state_round_trip(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    assert mod.load_state() is None  # yok
    mod.save_state(999, "no_change")
    assert mod.load_state()["last_successful_run_ms"] == 999


# ---------------------------------------------------------------------------
# Alan çıkarımı
# ---------------------------------------------------------------------------
def test_get_priority_dict_and_null(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    assert mod.get_priority({"priority": {"priority": "urgent", "color": "#f50000"}}) == "urgent"
    assert mod.get_priority({"priority": None}) == ""
    assert mod.get_priority({}) == ""


def test_get_description_present_and_absent(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    task = {"custom_fields": [{"id": mod.DESC_CF_ID, "type": "text", "value": "merhaba"}]}
    assert mod.get_description(task) == "merhaba"
    # value anahtarı yok → "" (gerçek API'de boş cf'de value absent)
    task2 = {"custom_fields": [{"id": mod.DESC_CF_ID, "type": "text"}]}
    assert mod.get_description(task2) == ""
    assert mod.get_description({"custom_fields": []}) == ""


def test_fetch_comment_summary_newest_first(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    # fetch_all_comments newest-first döner; comments[0] en yeni olmalı.
    fake = [
        {"id": "c2", "date": "200", "comment_text": "yeni"},
        {"id": "c1", "date": "100", "comment_text": "eski"},
    ]
    monkeypatch.setattr(mod, "fetch_all_comments", lambda s, tid: fake)
    latest, count = mod.fetch_comment_summary(None, "x")
    assert latest == "yeni"
    assert count == "2"


def test_fetch_comment_summary_empty_and_cap(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    monkeypatch.setattr(mod, "fetch_all_comments", lambda s, tid: [])
    assert mod.fetch_comment_summary(None, "x") == ("", "0")
    capped = [{"id": str(i), "date": str(i), "comment_text": "c"} for i in range(25)]
    monkeypatch.setattr(mod, "fetch_all_comments", lambda s, tid: capped)
    _, count = mod.fetch_comment_summary(None, "x")
    assert count == "25+"


# ---------------------------------------------------------------------------
# Ana akış: watermark sadece başarıda ilerler
# ---------------------------------------------------------------------------
FIXED_T0 = 9_000_000_000


def _wire_main(mod, monkeypatch, *, tasks, holiday=None, send_raises=False,
               fetch_raises=False):
    """main()'i ağ/SMTP olmadan çalıştırmak için modül-seviye fonksiyonları stub'la."""
    monkeypatch.setattr(mod, "now_ms", lambda: FIXED_T0)
    monkeypatch.setattr(mod, "today_istanbul", lambda: "GUNUMUZ")
    monkeypatch.setattr(mod, "turkish_holiday_name", lambda today: holiday)
    monkeypatch.setattr(mod, "get_session", lambda: object())
    if fetch_raises:
        def _raise_fetch(_s):
            raise RuntimeError("fetch eksik — watermark ilerlememeli")
        monkeypatch.setattr(mod, "fetch_all_tasks_strict", _raise_fetch)
    else:
        monkeypatch.setattr(mod, "fetch_all_tasks_strict", lambda s: tasks)
    # build_rows'u gerçek (yorum çekmeyen) haliyle; comment + excel'i stub'la
    monkeypatch.setattr(mod, "fetch_all_comments", lambda s, tid: [])
    monkeypatch.setattr(mod, "write_excel",
                        lambda rows: "Coalexus Pipeline Update_TEST.xlsx")
    monkeypatch.setattr(mod, "time", _NoSleep())

    sent = {"count": None, "recipients": None, "calls": 0}

    def fake_send(dosya, row_count, is_bootstrap, recipients=None):
        sent["calls"] += 1
        if send_raises:
            raise RuntimeError("SMTP patladı")
        sent["count"] = row_count
        sent["recipients"] = recipients

    monkeypatch.setattr(mod, "send_report_mail", fake_send)
    return sent


def _in_window_task():
    return {"id": "a", "name": "Yeni", "date_updated": str(FIXED_T0 - 100),
            "date_created": str(FIXED_T0 - 100), "due_date": None,
            "priority": None, "custom_fields": []}


class _NoSleep:
    """build_rows içindeki time.sleep'i no-op yapan minik shim."""

    def sleep(self, *_a, **_k):
        pass

    def time(self):
        return FIXED_T0 / 1000


def test_main_success_advances_watermark(monkeypatch, tmp_path):
    mod, state_path = _import_module(monkeypatch, tmp_path)
    W_old = 1_000
    mod.save_state(W_old, "sent")  # önceki başarılı çalışma
    tasks = [
        {"id": "a", "name": "Yeni", "date_updated": str(FIXED_T0 - 100),
         "date_created": str(FIXED_T0 - 100), "due_date": None, "priority": None,
         "custom_fields": []},
    ]
    sent = _wire_main(mod, monkeypatch, tasks=tasks)
    mod.main()
    assert sent["count"] == 1
    data = mod.load_state()
    assert data["last_successful_run_ms"] == FIXED_T0  # T0'a ilerledi
    assert data["last_status"] == "sent"


def test_watermark_not_advanced_on_send_failure(monkeypatch, tmp_path):
    """SMTP fail → state DEĞİŞMEZ.

    MUTATION: save_state, send_report_mail'den ÖNCE çağrılırsa (veya finally'de)
    bu test FAIL eder (watermark hatada ilerlerdi → boşluk).
    """
    mod, _ = _import_module(monkeypatch, tmp_path)
    W_old = 1_000
    mod.save_state(W_old, "sent")
    tasks = [
        {"id": "a", "name": "X", "date_updated": str(FIXED_T0), "date_created": "0",
         "due_date": None, "priority": None, "custom_fields": []},
    ]
    _wire_main(mod, monkeypatch, tasks=tasks, send_raises=True)
    with pytest.raises(RuntimeError):
        mod.main()
    assert mod.load_state()["last_successful_run_ms"] == W_old  # İLERLEMEDİ


def test_holiday_skips_mail_and_watermark(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    W_old = 1_000
    mod.save_state(W_old, "sent")
    sent = _wire_main(mod, monkeypatch, tasks=[], holiday="Kurban Bayramı")
    mod.main()
    assert sent["count"] is None  # mail çağrılmadı
    assert mod.load_state()["last_successful_run_ms"] == W_old  # İLERLEMEDİ


def test_no_change_day_sends_and_advances(monkeypatch, tmp_path):
    """0 task → kısa mail at + watermark ilerlet."""
    mod, _ = _import_module(monkeypatch, tmp_path)
    mod.save_state(1_000, "sent")
    # pencere dışı tek task → filtrelenip 0 kalsın
    tasks = [{"id": "old", "name": "Eski", "date_updated": "1", "date_created": "1",
              "due_date": None, "priority": None, "custom_fields": []}]
    sent = _wire_main(mod, monkeypatch, tasks=tasks)
    mod.main()
    assert sent["count"] == 0
    data = mod.load_state()
    assert data["last_successful_run_ms"] == FIXED_T0
    assert data["last_status"] == "no_change"


def test_main_fetch_failure_does_not_advance_watermark(monkeypatch, tmp_path):
    """Eksik/başarısız fetch → main fırlatır, watermark DEĞİŞMEZ (boşluk yok).

    Bu, 'kısmi/boş fetch'i tam pencere sanma' gap bug'ının regresyon testi.
    """
    mod, _ = _import_module(monkeypatch, tmp_path)
    mod.save_state(1_000, "sent")
    _wire_main(mod, monkeypatch, tasks=[], fetch_raises=True)
    with pytest.raises(RuntimeError):
        mod.main()
    assert mod.load_state()["last_successful_run_ms"] == 1_000  # İLERLEMEDİ


# ---------------------------------------------------------------------------
# Sıkı fetch: eksiklikte fırlatır, gerçekten boşsa [] döner
# ---------------------------------------------------------------------------
class _Resp:
    def __init__(self, status_code, payload=None):
        self.status_code = status_code
        self._payload = payload

    def json(self):
        if self._payload is None:
            raise ValueError("no json")
        return self._payload


def test_fetch_strict_raises_on_safe_get_none(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    monkeypatch.setattr(mod, "safe_get", lambda s, url, attempt_label="": None)
    with pytest.raises(RuntimeError):
        mod.fetch_all_tasks_strict(object())


def test_fetch_strict_raises_on_non_200(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    monkeypatch.setattr(mod, "safe_get", lambda s, url, attempt_label="": _Resp(500))
    with pytest.raises(RuntimeError):
        mod.fetch_all_tasks_strict(object())


def test_fetch_strict_raises_on_bad_json(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    monkeypatch.setattr(mod, "safe_get", lambda s, url, attempt_label="": _Resp(200, None))
    with pytest.raises(RuntimeError):
        mod.fetch_all_tasks_strict(object())


def test_fetch_strict_empty_list_is_ok(monkeypatch, tmp_path):
    """Gerçek bir 200 + boş liste meşru (boş gün); fırlatmaz."""
    mod, _ = _import_module(monkeypatch, tmp_path)
    monkeypatch.setattr(mod, "safe_get",
                        lambda s, url, attempt_label="": _Resp(200, {"tasks": []}))
    assert mod.fetch_all_tasks_strict(object()) == []


def test_fetch_strict_returns_tasks(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    payload = {"tasks": [{"id": "1"}, {"id": "2"}], "last_page": True}
    monkeypatch.setattr(mod, "safe_get",
                        lambda s, url, attempt_label="": _Resp(200, payload))
    out = mod.fetch_all_tasks_strict(object())
    assert [t["id"] for t in out] == ["1", "2"]


# ---------------------------------------------------------------------------
# Kısmi alıcı reddi → fırlat (sessiz kalıcı boşluk olmasın)
# ---------------------------------------------------------------------------
class _FakeSMTP:
    def __init__(self, refused):
        self._refused = refused

    def __call__(self, host, port, timeout=None):
        return self

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def login(self, u, p):
        pass

    def send_message(self, msg):
        return self._refused


def test_send_partial_refusal_raises(monkeypatch, tmp_path):
    import smtplib
    mod, _ = _import_module(monkeypatch, tmp_path)
    fake = _FakeSMTP({"efe.soykan@coalexus.com": (550, b"refused")})
    monkeypatch.setattr(mod.smtplib, "SMTP_SSL", fake)
    with pytest.raises(smtplib.SMTPRecipientsRefused):
        mod.send_report_mail(None, 0, False)


def test_send_all_accepted_ok(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    fake = _FakeSMTP({})  # boş dict = hepsi kabul
    monkeypatch.setattr(mod.smtplib, "SMTP_SSL", fake)
    mod.send_report_mail(None, 0, False)  # fırlatmamalı


# ---------------------------------------------------------------------------
# Test/override env'leri (DRY_RUN, RECIPIENT_OVERRIDE) — watermark ilerlememeli
# ---------------------------------------------------------------------------
def test_dry_run_skips_send_and_state(monkeypatch, tmp_path):
    """DRY_RUN → send çağrılmaz, state YAZILMAZ (mevcut watermark korunur)."""
    mod, _ = _import_module(monkeypatch, tmp_path)
    mod.save_state(1_000, "sent")
    monkeypatch.setenv("DRY_RUN", "1")
    sent = _wire_main(mod, monkeypatch, tasks=[_in_window_task()])
    mod.main()
    assert sent["calls"] == 0  # mail gönderilmedi
    assert mod.load_state()["last_successful_run_ms"] == 1_000  # İLERLEMEDİ


def test_recipient_override_sends_only_override_and_no_state(monkeypatch, tmp_path):
    """RECIPIENT_OVERRIDE → sadece override alıcılara gider; state YAZILMAZ."""
    mod, _ = _import_module(monkeypatch, tmp_path)
    mod.save_state(1_000, "sent")
    monkeypatch.setenv("RECIPIENT_OVERRIDE", "me@test.com, you@test.com")
    sent = _wire_main(mod, monkeypatch, tasks=[_in_window_task()])
    mod.main()
    assert sent["calls"] == 1
    assert sent["recipients"] == ["me@test.com", "you@test.com"]  # sabit 3 DEĞİL
    assert mod.load_state()["last_successful_run_ms"] == 1_000  # İLERLEMEDİ


def test_clean_run_advances_state(monkeypatch, tmp_path):
    """Kontrol: DRY_RUN/OVERRIDE YOKKEN watermark ilerler ve sabit alıcılara gider."""
    mod, _ = _import_module(monkeypatch, tmp_path)
    mod.save_state(1_000, "sent")
    monkeypatch.delenv("DRY_RUN", raising=False)
    monkeypatch.delenv("RECIPIENT_OVERRIDE", raising=False)
    sent = _wire_main(mod, monkeypatch, tasks=[_in_window_task()])
    mod.main()
    assert sent["recipients"] == mod.RECIPIENTS  # sabit 3 alıcı
    assert mod.load_state()["last_successful_run_ms"] == FIXED_T0  # İLERLEDİ


def test_parse_recipients(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    assert mod._parse_recipients("a@x.com, b@y.com ,, c@z.com") == [
        "a@x.com", "b@y.com", "c@z.com"]
    assert mod._parse_recipients("") == []
    assert mod._parse_recipients(None) == []


def test_write_excel_builds_file(monkeypatch, tmp_path):
    """Gerçek write_excel kapsaması (flow testleri write_excel'i stub'lar)."""
    import openpyxl
    mod, _ = _import_module(monkeypatch, tmp_path)
    monkeypatch.chdir(tmp_path)  # repo'yu kirletme
    rows = [["Ad", "ilk temas", "Açıklama", "Yorum", "01.01.2026",
             "01.01.2026 10:00", "urgent", "2"]]
    dosya = mod.write_excel(rows)
    assert os.path.exists(dosya)
    ws = openpyxl.load_workbook(dosya)["Pipeline"]
    assert [c.value for c in ws[1]] == mod.COLUMNS
    assert [c.value for c in ws[2]] == rows[0]


# ---------------------------------------------------------------------------
# Kolonlar: 8 kolon, "Aşama" 2. sırada; status çekimi; tam metin / rstrip
# ---------------------------------------------------------------------------
def test_columns_are_8_in_spec_order(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    assert mod.COLUMNS == [
        "Müşteri/Task Adı", "Aşama", "Açıklama", "Son Yorum",
        "Bitiş Tarihi", "Son Güncelleme", "Öncelik", "Yorum Sayısı",
    ]


def test_get_status_dict_string_missing(monkeypatch, tmp_path):
    """Ham /list/task status'u DICT döner; string'e normalize edilmiş hali de
    desteklenir; yoksa ''."""
    mod, _ = _import_module(monkeypatch, tmp_path)
    assert mod.get_status({"status": {"status": "ilk temas", "color": "#x"}}) == "ilk temas"
    assert mod.get_status({"status": "potansiyel alıcı"}) == "potansiyel alıcı"
    assert mod.get_status({}) == ""
    assert mod.get_status({"status": None}) == ""


def test_build_rows_8col_projection(monkeypatch, tmp_path):
    """build_rows 8 kolonu doğru sırayla üretir; Aşama=status, 2. kolonda."""
    mod, _ = _import_module(monkeypatch, tmp_path)
    monkeypatch.setattr(mod, "fetch_all_comments",
                        lambda s, tid: [{"id": "c", "date": "2", "comment_text": "son yorum"}])
    monkeypatch.setattr(mod, "time", _NoSleep())
    task = {
        "id": "t", "name": "Murat", "status": {"status": "ilk temas"},
        "due_date": "1785135600000", "date_updated": "1781708619365",
        "priority": {"priority": "urgent"},
        "custom_fields": [{"id": mod.DESC_CF_ID, "type": "text", "value": "kümes sahibi"}],
    }
    rows = mod.build_rows(None, [task])
    assert len(rows) == 1 and len(rows[0]) == 8
    r = rows[0]
    assert r[0] == "Murat"          # Müşteri/Task Adı
    assert r[1] == "ilk temas"      # Aşama
    assert r[2] == "kümes sahibi"   # Açıklama
    assert r[3] == "son yorum"      # Son Yorum
    assert r[6] == "urgent"         # Öncelik
    assert r[7] == "1"              # Yorum Sayısı


def test_clean_cell_rstrip_keeps_internal_newlines_and_caps(monkeypatch, tmp_path):
    mod, _ = _import_module(monkeypatch, tmp_path)
    assert mod._clean_cell("a\nb   \n\n") == "a\nb"   # iç \n kalır, son temizlenir
    assert mod._clean_cell("  x  ") == "  x"          # sadece rstrip (baş korunur)
    assert mod._clean_cell(None) == ""
    big = "y" * 40000
    assert len(mod._clean_cell(big)) == mod._EXCEL_CELL_MAX  # hard limit


def test_long_description_and_comment_not_truncated(monkeypatch, tmp_path):
    """500'lük cap KALDIRILDI: uzun Açıklama/Son Yorum tam yazılır (sadece
    rstrip)."""
    mod, _ = _import_module(monkeypatch, tmp_path)
    long_text = "k" * 5000
    task = {"custom_fields": [{"id": mod.DESC_CF_ID, "type": "text",
                              "value": long_text + "   \n\n"}]}
    desc = mod.get_description(task)
    assert desc == long_text  # tam metin, "…" yok, kelime ortasından kesik yok
    assert len(desc) == 5000

    monkeypatch.setattr(mod, "fetch_all_comments",
                        lambda s, tid: [{"id": "c", "date": "1",
                                        "comment_text": long_text + "\n  "}])
    latest, _count = mod.fetch_comment_summary(None, "x")
    assert latest == long_text
    assert len(latest) == 5000
